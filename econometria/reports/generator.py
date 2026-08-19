"""Geracao automatica de relatorios: HTML (visual, printavel), DOCX e XLSX.

Qualquer resultado de analise do sistema (regressao, series temporais, painel,
financas...) pode virar um relatorio: basta montar um ``ReportBuilder`` com
titulo, cartoes de indicadores (KPIs), secoes de texto/tabelas/graficos, e
exportar no formato desejado. Os textos narrativos (o "resumo_texto()" de cada
resultado) ja vem em portugues e com interpretacao pratica - o relatorio so
organiza isso visualmente.
"""
from __future__ import annotations

import base64
import io
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import pandas as pd
from jinja2 import Environment, FileSystemLoader, select_autoescape

_TEMPLATE_DIR = Path(__file__).parent / "templates"
_ENV = Environment(
    loader=FileSystemLoader(str(_TEMPLATE_DIR)),
    autoescape=select_autoescape(["html"]),
)


def fig_to_base64(fig, formato: str = "png", dpi: int = 140) -> str:
    """Converte uma figura matplotlib em uma string base64 pronta para embutir em HTML."""
    buf = io.BytesIO()
    fig.savefig(buf, format=formato, dpi=dpi, bbox_inches="tight")
    buf.seek(0)
    codificado = base64.b64encode(buf.read()).decode("ascii")
    return f"data:image/{formato};base64,{codificado}"


def _markdown_leve_para_html(texto: str) -> str:
    """Converte um subconjunto minimo de markdown (usado nos `resumo_texto()`) para HTML:
    **negrito**, quebras de paragrafo e listas com '- '.
    """
    blocos = texto.split("\n\n")
    html_blocos = []
    for bloco in blocos:
        bloco = bloco.strip()
        if not bloco:
            continue
        linhas = bloco.split("\n")
        if all(l.strip().startswith("- ") for l in linhas):
            itens_html = []
            for l in linhas:
                conteudo = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", l.strip()[2:])
                itens_html.append(f"<li>{conteudo}</li>")
            html_blocos.append(f"<ul>{''.join(itens_html)}</ul>")
        else:
            paragrafo = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", bloco)
            html_blocos.append(f"<p>{paragrafo}</p>")
    return "\n".join(html_blocos)


@dataclass
class ReportBuilder:
    titulo: str
    subtitulo: str = ""
    autor: str = ""
    _cards: list[dict] = field(default_factory=list, repr=False)
    _secoes: list[dict] = field(default_factory=list, repr=False)

    def add_card(self, rotulo: str, valor: str, dica: str = "") -> "ReportBuilder":
        """Adiciona um cartao de indicador (KPI) ao topo do relatorio."""
        self._cards.append({"rotulo": rotulo, "valor": valor, "dica": dica})
        return self

    def add_section(
        self,
        titulo: str,
        texto: str = "",
        tabelas: list[tuple[str, pd.DataFrame]] | None = None,
        figuras: list[tuple[str, str]] | None = None,
    ) -> "ReportBuilder":
        """Adiciona uma secao ao relatorio.

        Parameters
        ----------
        texto: texto narrativo (aceita '**negrito**' e listas com '- ').
        tabelas: lista de (legenda, DataFrame).
        figuras: lista de (legenda, data_uri) - use `fig_to_base64` para gerar o data_uri.
        """
        self._secoes.append(
            {
                "titulo": titulo,
                "texto_html": _markdown_leve_para_html(texto) if texto else "",
                "tabelas": [
                    {"legenda": legenda, "html": df.to_html(classes="tabela-dados", border=0, na_rep="-", float_format=lambda v: f"{v:,.4f}")}
                    for legenda, df in (tabelas or [])
                ],
                "figuras": [{"legenda": legenda, "src": src} for legenda, src in (figuras or [])],
            }
        )
        return self

    def render_html(self) -> str:
        template = _ENV.get_template("report.html.j2")
        return template.render(
            titulo=self.titulo,
            subtitulo=self.subtitulo,
            autor=self.autor,
            data_geracao=datetime.now().strftime("%d/%m/%Y as %H:%M"),
            cards=self._cards,
            secoes=self._secoes,
        )

    def save_html(self, caminho: str | Path) -> Path:
        caminho = Path(caminho)
        caminho.parent.mkdir(parents=True, exist_ok=True)
        caminho.write_text(self.render_html(), encoding="utf-8")
        return caminho

    def to_docx_bytes(self) -> bytes:
        """Gera o relatorio em formato Word (.docx) e retorna os bytes do arquivo."""
        from docx import Document
        from docx.shared import Pt, RGBColor

        doc = Document()
        titulo = doc.add_heading(self.titulo, level=0)
        if self.subtitulo:
            sub = doc.add_paragraph(self.subtitulo)
            sub.runs[0].italic = True
        doc.add_paragraph(f"Gerado em {datetime.now().strftime('%d/%m/%Y as %H:%M')}")

        if self._cards:
            doc.add_heading("Indicadores-chave", level=1)
            tabela = doc.add_table(rows=1, cols=len(self._cards))
            tabela.style = "Light Grid Accent 1"
            for cel, card in zip(tabela.rows[0].cells, self._cards):
                cel.text = f"{card['rotulo']}: {card['valor']}"

        for secao in self._secoes:
            doc.add_heading(secao["titulo"], level=1)
            texto_puro = re.sub(r"<[^>]+>", "", secao["texto_html"]).strip()
            if texto_puro:
                for paragrafo in texto_puro.split("\n"):
                    if paragrafo.strip():
                        doc.add_paragraph(paragrafo.strip())
            for tab in secao["tabelas"]:
                doc.add_paragraph(tab["legenda"]).runs[0].bold = True
                df_tables = pd.read_html(io.StringIO(tab["html"]))
                if df_tables:
                    df = df_tables[0]
                    tabela_doc = doc.add_table(rows=1, cols=len(df.columns))
                    tabela_doc.style = "Light List Accent 1"
                    for i, col in enumerate(df.columns):
                        tabela_doc.rows[0].cells[i].text = str(col)
                    for _, row in df.iterrows():
                        cells = tabela_doc.add_row().cells
                        for i, val in enumerate(row):
                            cells[i].text = str(val)

        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    def to_excel_bytes(self, tabelas_extra: dict[str, pd.DataFrame] | None = None) -> bytes:
        """Exporta todas as tabelas do relatorio (mais quaisquer `tabelas_extra`) para um .xlsx com varias abas."""
        buf = io.BytesIO()
        with pd.ExcelWriter(buf, engine="xlsxwriter") as writer:
            aba_usada = set()

            def nome_unico(base: str) -> str:
                nome = re.sub(r"[\\/*?:\[\]]", "", base)[:31] or "Tabela"
                i = 1
                candidato = nome
                while candidato in aba_usada:
                    sufixo = f"_{i}"
                    candidato = nome[: 31 - len(sufixo)] + sufixo
                    i += 1
                aba_usada.add(candidato)
                return candidato

            for secao in self._secoes:
                for tab in secao["tabelas"]:
                    df_tables = pd.read_html(io.StringIO(tab["html"]))
                    if df_tables:
                        df_tables[0].to_excel(writer, sheet_name=nome_unico(tab["legenda"] or secao["titulo"]))
            for nome, df in (tabelas_extra or {}).items():
                df.to_excel(writer, sheet_name=nome_unico(nome))
        return buf.getvalue()
